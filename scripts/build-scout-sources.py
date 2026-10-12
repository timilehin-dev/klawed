#!/usr/bin/env python
"""Build klaw/docs/scout-sources.md from the Gemini source-directory docx.

Consolidates the 15 source tables + search dictionary into the Scout's hunting
config, adds studentoffers.co, and appends the operational notes (heuristics,
schedule, traps) extracted from the document narrative.
"""
import docx, re, os

DOCX = r"C:\Users\HP\Downloads\Builder Opportunity Source Directory.docx"
OUT = r"C:\Users\HP\klaw\docs\scout-sources.md"

d = docx.Document(DOCX)

# table index -> lane label
LANE_MAP = {
    0: "Jobs & Remote",
    1: "Hackathons, Competitions & Bounties",
    2: "Grants & Non-Dilutive Funding",
    3: "Accelerators & Startup Studios",
    4: "Visas & Talent Immigration",
    5: "Scholarships & Fellowships",
    6: "AI Credits & Student Programmes",
    7: "Media Desks & Opportunity Feeds",
    8: "Social — X (Twitter)",
    9: "Social — LinkedIn",
    10: "Social — Facebook",
    11: "Social — Telegram & WhatsApp",
    12: "Social — Discord",
    13: "Social — Reddit & Forums",
    14: "Newsletters & Digests",
}

def cell(r, i):
    return r.cells[i].text.strip()

def emit_table(tb, skip_first_col_label=None):
    hdr = [cell(tb.rows[0], i) for i in range(len(tb.columns))]
    lines = []
    ncols = len(tb.columns)
    # Social tables (T8 X, T13 Reddit) have 6 cols and NO "Class" column:
    # Name | URL/Handle | Lanes | Region | Cadence | Note
    # Insert an empty Class so every row aligns to the 7-col schema.
    has_class = any("class" in h.lower() or "platform" in h.lower() or "type" in h.lower()
                    or "delivery" in h.lower() for h in hdr)
    for r in tb.rows[1:]:
        vals = [cell(r, i) for i in range(ncols)]
        if ncols == 6 and not has_class:
            vals = [vals[0], vals[1], "—", vals[2], vals[3], vals[4], vals[5]]
        lines.append(vals)
    return hdr, lines

parts = []
parts.append("""# Klawed — Scout Hunting Sources

> The Scout's target list. Generated from the Gemini source directory
> (`Builder Opportunity Source Directory.docx`) + `studentoffers.co`.
> **213 sources** across every board lane. Hunt these DIRECTLY — do not fall back
> to generic web searches, which surface expired/recycled content.
>
> Each entry: **Name | URL | Class | Lanes | Region | Cadence | What to look for**
> `[verify]` = unconfirmed by the source; the Verifier must check it before use.

## How the Scout uses this file
1. Work lane by lane. For each lane, fetch/extract from the listed canonical sources.
2. Apply the **Search-Term Dictionary** (§Search) for phrase-level detection.
3. Apply the **Freshness Heuristics** (§Freshness) before accepting anything.
4. Honour the **Ingestion Schedule** (§Schedule) — daily sources every run, weekly
   sources on the weekly run, per-cohort sources when their cycle is near.
5. Everything found goes to `quarantine/raw/scout_<date>.md` in the standard format.

""")

# --- source tables ---
for idx, label in LANE_MAP.items():
    tb = d.tables[idx]
    hdr, rows = emit_table(tb)
    parts.append(f"\n## {label}\n")
    if idx == 6:
        parts.append("_Includes `studentoffers.co` (596+ verified student discounts, "
                     "category-indexed, regularly refreshed) — a strong cross-check for "
                     "the student_deals lane._\n")
    parts.append("| Source | URL | Class | Lanes | Region | Cadence | What to look for |")
    parts.append("|---|---|---|---|---|---|---|")
    for v in rows:
        # pad to 7 cols
        v = (v + [""] * 7)[:7]
        name, url, cls, lanes, region, cad, note = v
        note = note.replace("|", "/").replace("\n", " ")
        parts.append(f"| {name} | {url} | {cls} | {lanes} | {region} | {cad} | {note} |")
    # add studentoffers.co into the AI credits / student programmes lane
    if idx == 6:
        parts.append("| StudentOffers.co | https://www.studentoffers.co/tools | Verified "
                     "student-discount directory (596+ offers, category-indexed, refreshed "
                     "regularly) | student_deals, ai_credits | Global | Weekly | Category "
                     "index at /tools?category=… — ai-and-machine-learning, cloud-and-hosting, "
                     "development-tools, productivity, design-and-creative; each offer has a "
                     "direct claim link. Also carries the GitHub Student Pack as a featured "
                     "pack. Cross-check against education.github.com/pack. |")

# --- search dictionary (T15) ---
tb = d.tables[15]
parts.append("\n## §Search — Per-Category Search-Term Dictionary\n")
parts.append("Use these EXACT high-intent phrases. Never search bare words like "
             "\"job\" or \"hackathon\".\n")
parts.append("| Lane | Primary intake strings | Freshness/boundary markers | Exclusions |")
parts.append("|---|---|---|---|")
for r in tb.rows[1:]:
    v = [cell(r, i) for i in range(len(r.cells))]
    v = (v + [""] * 4)[:4]
    parts.append("| " + " | ".join(x.replace("|", "/").replace("\n", " ") for x in v) + " |")

# --- operational notes from narrative ---
parts.append("""
## §Freshness — Freshness Heuristics (apply before accepting any find)

1. **HTTP status + DOM introspection.** Drop 404 / 410 / 502 / redirect loops >3 hops.
   But a 200 can still be a CLOSED role — parse the DOM for closure strings:
   Greenhouse: "This job is no longer available" / missing `#application-form`;
   Lever: "This position has been closed" / missing `.postings-btn-wrapper`;
   Ashby: "This job is no longer accepting applications";
   Workable: banner "This vacancy has closed". A closed-state page is a DROP.
2. **Semantic date extraction + sliding window.** Parse any explicit deadline
   ("Applications close:", "Apply before:"). Past date = DROP. For rolling items,
   take the publish timestamp; if `now - published > 60 days` = DROP as stale,
   UNLESS the page carries an override ("rolling basis", "continuous review").
   Evergreen pages with no publication metadata are suspect — treat with caution.
3. **Application mechanism must be interactive.** There must be a live form / apply
   widget (ATS form, Airtable, Typeform, Google Form). Follow hosted forms and check
   they are not closed ("The form is no longer accepting responses" = DROP).
4. **Canonical link resolution + dedup.** Unroll every shortener (bit.ly, t.co,
   ow.ly) to the final landing page BEFORE verifying. Dedup key:
   `SHA256(normalized org + normalized title + normalized location)` — a match
   within the preceding 90 days means update-or-discard, never double-post.
5. **Platform-specific status.** DoraHacks/Devpost expose JSON-LD status fields —
   read them directly. For scholarships, validate the academic cycle: notices for
   past years (2023/2024, 2024/2025) are DROPPED.

## §Schedule — Ingestion Cadence

**DAILY** (high-yield, listings fill in hours):
- Corporate ATS: Moniepoint (Workable), Flutterwave (Greenhouse), Paystack (Lever),
  Deel (Ashby), Remote.com (Ashby)
- Telegram: t.me/jobsregion, t.me/scholarshipregion, t.me/jbtoday
- X curators: @ChuksEmma947401, @RemotejobWtessy, @Opps_DeskHQ
- Aggregator APIs: himalayas.app/jobs, We Work Remotely RSS (programming),
  Devpost online-competitions queue

**WEEKLY** (Sunday evening / Monday morning):
- Media: Techpoint Africa Jobs Weekly, TechCabal Entering Tech, Condia
- Hackathon/bounty: Zindi competitions, DoraHacks Africa/AI, Superteam Earn Nigeria
- Communities: @dev_careers (DevCareer), @SheCodeAfrica, OSCA Discord #announcements,
  Web3 Nigeria Discord
- Aggregators: opportunitiesforafricans.com, Opportunity Desk, studentoffers.co

**PER-COHORT / ON-DEMAND** (14–30 day cycles or fixed calendar triggers):
- Accelerators: Y Combinator (bi-annual), ARM Labs Techstars Lagos, Google for
  Startups Accelerator Africa, The Baobab Network
- Corporate hackathons: Wema Hackaholics, Superteam Nigeria Superhack
- Scholarships: Chevening (Aug–Nov window), Mastercard Foundation partner pages,
  DAAD EP.1019, Gates Cambridge
- Immigration: UK Global Talent (Tech Nation), Germany Chancenkarte, Portugal Vistos
  MNE digital nomad, Japan MOFA specified-activity visas

## §Top20 — Highest-value sources for a Nigeria-first audience

Paystack Careers (Lever) · Moniepoint Careers (Workable) · Flutterwave Careers
(Greenhouse) · Jobs Region Telegram · Scholarship Region Telegram · Wema Hackaholics ·
Superteam Nigeria / superhack.fun · Open Source Community Africa Discord · DevCareer
(@dev_careers) · She Code Africa (@SheCodeAfrica) · The Baobab Network · ARM Labs
Lagos Techstars · Techpoint Africa Jobs Weekly · Condia · TechCabal Entering Tech ·
Himalayas Remote · Opportunities for Africans · GitHub Student Developer Pack ·
Make it in Germany (Chancenkarte) · UK Global Talent (Tech Nation)

## §Traps — Structural pathologies & fraud to actively filter

1. **Zombie aggregator feedback loop.** Low-tier sites mirror each other's scrapes
   recursively, so a listing removed from the primary ATS can persist across dozens
   of secondary sites for months. ALWAYS resolve to the primary ATS page — never
   trust a repost as proof the role is live.
2. **Engagement-farming social accounts.** Curators rewrite OLD/closed listings to
   grow followers ("Drop your CV", "Tag a friend", "Repost"). Treat undated social
   reposts as unverified until the primary source confirms.
3. **Evergreen "zombie" portals.** Some banks/telcos keep "Talent Pool" /
   "Spontaneous Application" forms open forever. These are NOT live vacancies —
   a permanent application form is not an open position. DROP unless a specific
   role is named.
4. **Fee scams.** Fake recruiters/embassies/scholarship bodies demanding payment
   for "processing" a remote role or relocation. Any request for money = DROP.
5. **The "global remote" filtering trap.** Boards tag listings "Remote" while the
   body restricts to US/EU residency, US citizenship, or EST-timezone overlap.
   Read the eligibility text, not the tag. Tag region honestly (`nigeria`/`africa`/
   `global`) based on what a Nigerian builder can actually apply for.

## §Internships note
No dedicated internships table exists in the source directory — internships ride
inside the jobs/ATS tables (most corporate ATS boards above carry `internships` in
their lanes) and surface through the media and Telegram/community lanes. Hunt
internships through those sources with the `internship` lane strings.

## §Verify list
Entries marked `[verify]` in the tables above (e.g. NaijaHacks) must be link-checked
by the Verifier before any listing from them is published.
""")

os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(parts))

size = os.path.getsize(OUT)
print(f"WROTE {OUT} ({size} bytes)")
